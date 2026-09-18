from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from . import powerpoint
from .cancellation import CancellationRequested

ROOT = Path(__file__).parents[1]


def resolve_ffmpeg(project_root: Path | None = None) -> str:
    candidates = [
        os.getenv("FFMPEG_PATH"),
        shutil.which("ffmpeg"),
        shutil.which("ffmpeg.exe"),
        str((project_root or ROOT) / ".tools" / "ffmpeg" / "ffmpeg.exe"),
        str(ROOT / ".tools" / "ffmpeg" / "ffmpeg.exe"),
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(Path(candidate).resolve())
    raise RuntimeError("FFmpeg was not found. Set FFMPEG_PATH or install it at <project>/.tools/ffmpeg/ffmpeg.exe.")


def command(name: str) -> str:
    path = shutil.which(name)
    if not path:
        raise RuntimeError(f"Required executable not found: {name}")
    return path


def run(args: list[str], cancel_check=None) -> None:
    process = subprocess.Popen(args)
    while process.poll() is None:
        if cancel_check and cancel_check():
            process.terminate()
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait()
            raise CancellationRequested("Operation cancelled by the user")
        try:
            process.wait(timeout=0.25)
        except subprocess.TimeoutExpired:
            continue
    if process.returncode:
        raise subprocess.CalledProcessError(process.returncode, args)


def _video_from_slide_images(
    slides: list[Path],
    temp: Path,
    visual: Path,
    duration: float,
    gap_seconds: float,
    ffmpeg: str,
    cancel_check=None,
) -> Path:
    if not slides:
        raise RuntimeError("No rendered slides found")
    per_slide = duration / len(slides)
    concat = temp / "slides.txt"
    with concat.open("w", encoding="utf-8") as file:
        for slide in slides:
            file.write(f"file '{slide.resolve()}'\n")
            file.write(f"duration {per_slide + gap_seconds:.3f}\n")
        file.write(f"file '{slides[-1].resolve()}'\n")
    run([ffmpeg, "-y", "-f", "concat", "-safe", "0", "-i", str(concat), "-r", "30", "-c:v", "libx264", "-pix_fmt", "yuv420p", str(visual)], cancel_check)
    return visual


def _render_pdf_with_pymupdf(
    pdf: Path,
    temp: Path,
    visual: Path,
    duration: float,
    gap_seconds: float,
    ffmpeg: str,
    cancel_check=None,
) -> Path:
    """Rasterize a PowerPoint-generated PDF without requiring Poppler."""
    import fitz

    slides_dir = temp / "powerpoint-slides"
    slides_dir.mkdir(parents=True, exist_ok=True)
    slides: list[Path] = []
    with fitz.open(str(pdf)) as document:
        for index, page in enumerate(document, start=1):
            if cancel_check and cancel_check():
                raise CancellationRequested("Operation cancelled by the user")
            image = slides_dir / f"slide-{index:04d}.png"
            page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(str(image))
            slides.append(image)
    return _video_from_slide_images(slides, temp, visual, duration, gap_seconds, ffmpeg, cancel_check)


def audio_duration(path: Path, ffmpeg: str) -> float:
    result = subprocess.run([ffmpeg, "-i", str(path)], capture_output=True, text=True, check=False)
    import re
    match = re.search(r"Duration: (\d+):(\d+):(\d+\.\d+)", result.stderr)
    if not match:
        raise RuntimeError(f"Could not read audio duration: {path}")
    h, m, s = match.groups()
    return int(h) * 3600 + int(m) * 60 + float(s)


def _fit_audio_filter(source_duration: float, target_duration: float) -> str:
    if target_duration >= source_duration:
        padding = target_duration - source_duration
        return f"apad=pad_dur={padding:.3f},atrim=duration={target_duration:.3f}"
    return _tempo_filter(source_duration, target_duration)


def _tempo_filter(source_duration: float, target_duration: float) -> str:
    ratio = source_duration / target_duration
    filters: list[str] = []
    while ratio < 0.5:
        filters.append("atempo=0.5")
        ratio /= 0.5
    while ratio > 2.0:
        filters.append("atempo=2.0")
        ratio /= 2.0
    filters.append(f"atempo={ratio:.6f}")
    return ",".join(filters)


def prepare_review_audio(primary: Path, secondary: Path, output_dir: Path, target_duration_seconds: float | None, cancel_check=None) -> tuple[Path, Path, bool]:
    """Create the exact audio files whose timeline is used by review subtitles."""
    output_dir.mkdir(parents=True, exist_ok=True)
    primary_review = output_dir / "review-primary.wav"
    secondary_review = output_dir / "review-secondary.wav"
    if not target_duration_seconds:
        shutil.copy2(primary, primary_review)
        shutil.copy2(secondary, secondary_review)
        return primary_review, secondary_review, False
    try:
        ffmpeg = resolve_ffmpeg(output_dir.parent)
    except RuntimeError:
        # Audio can still be reviewed without normalization; video rendering
        # will report the missing FFmpeg dependency separately.
        shutil.copy2(primary, primary_review)
        shutil.copy2(secondary, secondary_review)
        return primary_review, secondary_review, False
    source_duration = max(audio_duration(primary, ffmpeg), audio_duration(secondary, ffmpeg))
    audio_filter = _fit_audio_filter(source_duration, target_duration_seconds)
    for source, destination in ((primary, primary_review), (secondary, secondary_review)):
        run([ffmpeg, "-y", "-i", str(source), "-filter:a", audio_filter, "-ar", "24000", "-ac", "1", "-c:a", "pcm_s16le", str(destination)], cancel_check)
    return primary_review, secondary_review, True


def render_presentation_pdf(presentation: Path, output_dir: Path, cancel_check=None) -> Path:
    """Render the reviewed PPT to PDF before speech/video generation."""
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf = output_dir / "presentation.pdf"
    with tempfile.TemporaryDirectory(prefix="ai-video-pdf-") as temp_name:
        temp = Path(temp_name)
        renderer = os.getenv("VIDEO_RENDERER", "auto").lower()
        use_powerpoint = renderer == "powerpoint" or (renderer == "auto" and powerpoint.available())
        if renderer not in {"auto", "powerpoint", "libreoffice"}:
            raise ValueError("VIDEO_RENDERER must be auto, powerpoint, or libreoffice")
        if use_powerpoint:
            from pptx import Presentation
            slide_count = len(Presentation(str(presentation)).slides)
            try:
                powerpoint.render(presentation, pdf, temp / "visual.mp4", [1.0] * slide_count)
            except RuntimeError:
                # In auto mode, a detected Mac PowerPoint may still reject
                # AppleScript automation because of an Office/version or
                # Automation-permission issue. Keep the cross-platform fallback
                # usable; explicit VIDEO_RENDERER=powerpoint still fails loudly.
                if renderer == "powerpoint":
                    raise
                use_powerpoint = False
        if not use_powerpoint:
            soffice = command("soffice")
            run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(temp), str(presentation)], cancel_check)
            converted = temp / f"{presentation.stem}.pdf"
            if not converted.exists():
                raise RuntimeError(f"LibreOffice did not create {converted}")
            shutil.copy2(converted, pdf)
    return pdf


def build_video(
    presentation: Path,
    primary_audio: Path,
    secondary_audio: Path,
    output_dir: Path,
    gap_seconds: float,
    target_duration_seconds: float | None = None,
    primary_language: str = "en-US",
    secondary_language: str = "zh-CN",
    cancel_check=None,
) -> dict[str, Any]:
    ffmpeg = resolve_ffmpeg(output_dir.parent)
    renderer = os.getenv("VIDEO_RENDERER", "auto").lower()
    if renderer not in {"auto", "powerpoint", "libreoffice"}:
        raise ValueError("VIDEO_RENDERER must be auto, powerpoint, or libreoffice")
    use_powerpoint = renderer == "powerpoint" or (renderer == "auto" and powerpoint.available())
    soffice = pdftoppm = None
    if not use_powerpoint:
        soffice = command("soffice")
        pdftoppm = command("pdftoppm")
    output_dir.mkdir(parents=True, exist_ok=True)
    slide_count = 0
    with tempfile.TemporaryDirectory(prefix="ai-video-") as temp_name:
        temp = Path(temp_name)
        primary_source = primary_audio
        secondary_source = secondary_audio
        source_duration = max(audio_duration(primary_audio, ffmpeg), audio_duration(secondary_audio, ffmpeg))
        target_duration = target_duration_seconds or source_duration
        if target_duration_seconds:
            primary_source = temp / "primary-normalized.wav"
            secondary_source = temp / "secondary-normalized.wav"
            for source, destination in ((primary_audio, primary_source), (secondary_audio, secondary_source)):
                run([ffmpeg, "-y", "-i", str(source), "-filter:a", _fit_audio_filter(source_duration, target_duration), "-ar", "24000", "-ac", "1", "-c:a", "pcm_s16le", str(destination)], cancel_check)
        duration = target_duration
        visual = temp / "visual.mp4"
        if use_powerpoint and powerpoint.is_macos():
            from pptx import Presentation
            slide_count = len(Presentation(str(presentation)).slides)
            # PowerPoint for Mac can export a faithful PDF, but does not
            # expose the Windows COM CreateVideo API. Render that PDF with
            # the already-required PyMuPDF dependency instead of Poppler.
            pdf = temp / "presentation.pdf"
            try:
                powerpoint.export_pdf_macos(presentation, pdf)
            except RuntimeError:
                # Auto detection should not make video generation fail when
                # AppleScript is unavailable or PowerPoint rejects the deck.
                if renderer == "powerpoint":
                    raise
                use_powerpoint = False
            else:
                shutil.copy2(pdf, output_dir / "presentation.pdf")
                _render_pdf_with_pymupdf(pdf, temp, visual, duration, gap_seconds, ffmpeg, cancel_check)
        if use_powerpoint and not powerpoint.is_macos():
            from pptx import Presentation
            slide_count = len(Presentation(str(presentation)).slides)
            per_slide = duration / slide_count
            visual_ppt = temp / "visual.pptx"
            powerpoint.render(presentation, temp / "presentation.pdf", visual_ppt, [per_slide + gap_seconds] * slide_count)
            shutil.copy2(temp / "presentation.pdf", output_dir / "presentation.pdf")
        elif not use_powerpoint:
            soffice = command("soffice")
            pdftoppm = command("pdftoppm")
            run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(temp), str(presentation)], cancel_check)
            pdf = temp / f"{presentation.stem}.pdf"
            if not pdf.exists():
                raise RuntimeError(f"LibreOffice did not create {pdf}")
            shutil.copy2(pdf, output_dir / "presentation.pdf")
            run([pdftoppm, "-png", "-r", "144", str(pdf), str(temp / "slide")], cancel_check)
            slides = sorted(temp.glob("slide-*.png"))
            visual = temp / "visual.mp4"
            _video_from_slide_images(slides, temp, visual, duration, gap_seconds, ffmpeg, cancel_check)
        dual = output_dir / "presentation-dual.mp4"
        run([ffmpeg, "-y", "-i", str(visual), "-i", str(primary_source), "-i", str(secondary_source), "-map", "0:v:0", "-map", "1:a:0", "-map", "2:a:0", "-c:v", "copy", "-c:a", "aac", "-b:a", "128k", "-metadata:s:a:0", f"language={primary_language}", "-metadata:s:a:1", f"language={secondary_language}", "-shortest", str(dual)], cancel_check)
        for name, audio, language in (("presentation-primary.mp4", primary_source, primary_language), ("presentation-secondary.mp4", secondary_source, secondary_language)): 
            run([ffmpeg, "-y", "-i", str(visual), "-i", str(audio), "-map", "0:v:0", "-map", "1:a:0", "-c:v", "copy", "-c:a", "copy", "-metadata:s:a:0", f"language={language}", "-shortest", str(output_dir / name)], cancel_check)
    return {"output_dir": str(output_dir), "slides": slide_count, "duration_seconds": duration}
